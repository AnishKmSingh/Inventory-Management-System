import ttkbootstrap as tb
from ttkbootstrap.constants import *
from tkinter import TclError, messagebox, ttk
from backend import Database
import tkinter as tk
from tkinter import filedialog
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

db = Database()

FONT_TITLE = ("Segoe UI", 20, "bold")
FONT_HEADING = ("Segoe UI", 14, "bold")
FONT_NORMAL = ("Segoe UI", 11)
FONT_BIG = ("Segoe UI", 26, "bold")

class LoginWindow(tb.Window):
    PURPLE = "#6d28d9"
    PURPLE_DARK = "#5b21b6"
    GOLD = "#fbbf24"

    def __init__(self):
        super().__init__(themename="pulse")
        self.dark_mode = False
        self.current_page = "show_dashboard"
        self.title("Inventory Management System - Login")
        self.geometry("1000x620")
        self.minsize(900, 600)

        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=3, uniform="half")
        self.grid_columnconfigure(1, weight=2, uniform="half")

        # ---------- LEFT: gradient hero panel ----------
        self.hero = tk.Canvas(self, highlightthickness=0, bd=0)
        self.hero.grid(row=0, column=0, sticky="nsew")
        self.hero.bind("<Configure>", self._draw_hero)

        # ---------- RIGHT: login card ----------
        right = tk.Frame(self, bg="#f3f4f8")
        right.grid(row=0, column=1, sticky="nsew")

        card = tk.Frame(right, bg="white", padx=40, pady=40,
                        highlightthickness=1, highlightbackground="#e3e5ee")
        card.place(relx=0.5, rely=0.5, anchor="center")

        tk.Label(card, text="Welcome Back", bg="white", fg="#1e1b4b",
                 font=("Segoe UI", 22, "bold")).pack(anchor="w")
        tk.Label(card, text="Sign in to manage your inventory", bg="white",
                 fg="#6b7280", font=("Segoe UI", 10)).pack(anchor="w", pady=(2, 25))

        def field(label, show=""):
            tk.Label(card, text=label, bg="white", fg="#6b7280",
                     font=("Segoe UI", 9, "bold")).pack(anchor="w")
            entry = tk.Entry(card, width=30, show=show, font=("Segoe UI", 11),
                             relief="flat", bg="#f3f4f8", fg="#111827",
                             insertbackground="#111827", highlightthickness=1,
                             highlightbackground="#d9dbe5",
                             highlightcolor=self.PURPLE)
            entry.pack(fill="x", pady=(4, 16), ipady=8)
            return entry

        def button(text, bg, fg, hover, command, border=None):
            b = tk.Button(card, text=text, command=command, bg=bg, fg=fg,
                          activebackground=hover, activeforeground=fg,
                          relief="flat", bd=0, cursor="hand2",
                          font=("Segoe UI", 11, "bold"),
                          highlightthickness=1 if border else 0,
                          highlightbackground=border or bg)
            b.bind("<Enter>", lambda e: b.config(bg=hover))
            b.bind("<Leave>", lambda e: b.config(bg=bg))
            return b

        self.username_entry = field("USERNAME")
        self.password_entry = field("PASSWORD", show="•")
        self.password_entry.bind("<Return>", lambda e: self.login())

        button("Login", self.PURPLE, "white", self.PURPLE_DARK,
               self.login).pack(fill="x", ipady=8, pady=(6, 0))
        button("Create Account", "white", self.PURPLE, "#ede9fe",
               self.open_signup, border=self.PURPLE).pack(fill="x", ipady=8, pady=(10, 18))

        tk.Label(card, text="Default: admin / admin@123", bg="white",
                 fg="#9ca3af", font=("Segoe UI", 8)).pack()

    @staticmethod
    def _blend(c1, c2, t):
        a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
        b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
        return "#%02x%02x%02x" % tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3))

    def _draw_hero(self, event=None):
        c = self.hero
        w, h = c.winfo_width(), c.winfo_height()
        c.delete("all")

        # gradient background
        for y in range(0, h, 2):
            c.create_rectangle(0, y, w, y + 2, outline="",
                               fill=self._blend("#5b21b6", "#1e1b4b", y / h))

        # decorative rings
        c.create_oval(w - 280, -140, w + 120, 260, outline="#7c4dd8", width=2)
        c.create_oval(w - 210, -70, w + 50, 190, outline="#6d3fc9", width=2)
        c.create_oval(-160, h - 220, 200, h + 140, outline="#7c4dd8", width=2)

        x = max(50, int(w * 0.12))
        y = max(40, h // 2 - 255)

        c.create_text(x, y, text="📦", font=("Segoe UI Emoji", 36),
                      anchor="w", fill="white")
        c.create_text(x, y + 75, text="Inventory Management System",
                      font=("Segoe UI", 25, "bold"), anchor="w", fill="white")
        c.create_rectangle(x, y + 105, x + 60, y + 109, fill=self.GOLD, outline="")
        c.create_text(x, y + 135,
                      text="Track stock, orders and suppliers in one place.",
                      font=("Segoe UI", 12), anchor="w", fill="#d8ccf5")

        c.create_line(x, y + 170, w - x, y + 170, fill="#8b6fd6")

        features = (
            ("📊", "Live dashboard with low-stock alerts"),
            ("🚚", "Purchase and sales order tracking"),
            ("📍", "Multi-location storage and stock transfers"),
            ("📜", "Product history and full audit log"),
        )
        for i, (icon, text) in enumerate(features):
            fy = y + 205 + i * 42
            c.create_oval(x, fy - 15, x + 30, fy + 15,
                        fill="#7c4dd8", outline=self.GOLD, width=1)
            c.create_text(x + 15, fy, text=icon, font=("Segoe UI Emoji", 13),
                        fill="white")
            c.create_text(x + 46, fy, text=text, font=("Segoe UI", 12),
                        anchor="w", fill="#f3effd")

            c.create_line(x, y + 370, w - x, y + 370, fill="#8b6fd6")

                # ---------- credits card ----------
        card_h = 78
        top = max(y + 395, h - card_h - 30)
        c.create_rectangle(x, top, w - x, top + card_h,
                           fill="#4a2a99", outline="#7c4dd8")
        c.create_rectangle(x, top, x + 4, top + card_h, fill=self.GOLD, outline="")

        for i, letter in enumerate(("A", "R")):
            ax = x + 24 + i * 34
            c.create_oval(ax, top + 20, ax + 38, top + 58,
                          fill="#7c4dd8", outline=self.GOLD, width=2)
            c.create_text(ax + 19, top + 39, text=letter,
                          font=("Segoe UI", 13, "bold"), fill="white")

        tx = x + 24 + 34 + 38 + 16
        c.create_text(tx, top + 18, text="DEVELOPED BY", anchor="w",
                      font=("Segoe UI", 8, "bold"), fill=self.GOLD)
        c.create_text(tx, top + 38, text="Anish Kumar Singh & Rudra Bounthiyal",
                      anchor="w", font=("Segoe UI", 11, "bold"), fill="#ffffff")
        c.create_text(tx, top + 59,
                      text="TECH STACK -> Python Tkinter + MySQL",
                      anchor="w", font=("Segoe UI", 9), fill="#cfc3f0")
    def login(self):
        username = self.username_entry.get().strip()
        password = self.password_entry.get()
        if not username or not password:
            messagebox.showwarning(
                "Login Required",
                "Enter both your username and password."
            )
            return

        user = db.authenticate(username, password)
        if user:
            self.destroy()
            MainApp(user).mainloop()
        else:
            messagebox.showerror("Login Failed", "Invalid username or password")

    def open_signup(self):
        signup = tb.Toplevel(self)
        signup.title("Create Account")
        signup.geometry("380x470")
        signup.resizable(False, False)
        signup.transient(self)
        signup.grab_set()

        card = tb.Frame(signup, padding=30)
        card.pack(fill=BOTH, expand=True)
        tb.Label(card, text="Create Account", font=FONT_HEADING).pack(pady=(0, 15))

        tb.Label(card, text="Username", font=FONT_NORMAL).pack(anchor=W)
        username_entry = tb.Entry(card, bootstyle="primary")
        username_entry.pack(fill=X, pady=(2, 10), ipady=5)

        tb.Label(card, text="Password", font=FONT_NORMAL).pack(anchor=W)
        password_entry = tb.Entry(card, show="*", bootstyle="primary")
        password_entry.pack(fill=X, pady=(2, 10), ipady=5)

        tb.Label(card, text="Confirm Password", font=FONT_NORMAL).pack(anchor=W)
        confirm_entry = tb.Entry(card, show="*", bootstyle="primary")
        confirm_entry.pack(fill=X, pady=(2, 10), ipady=5)

        def register():
            username = username_entry.get().strip()
            password = password_entry.get()
            if password != confirm_entry.get():
                messagebox.showerror(
                    "Sign Up Failed", "Passwords do not match", parent=signup
                )
                return
            try:
                db.create_user(username, password)
            except ValueError as error:
                messagebox.showerror("Sign Up Failed", str(error), parent=signup)
                return
            messagebox.showinfo(
                "Account Created",
                "Your account was created. You can now log in.",
                parent=signup
            )
            signup.destroy()
            self.username_entry.delete(0, END)
            self.username_entry.insert(0, username)
            self.password_entry.focus_set()

        tb.Button(
            card, text="Create Account", bootstyle=SUCCESS,
            command=register
        ).pack(fill=X, ipady=7)
        confirm_entry.bind("<Return>", lambda event: register())
        username_entry.focus_set()


class MainApp(tb.Window):
    NAV_ITEMS = [
        ("🏠  Dashboard", "show_dashboard"),
        ("📦  Products", "show_products"),
        ("🚚  Suppliers", "show_suppliers"),
        ("👥  Customers", "show_customers"),
        ("📍  Storage Locations", "show_storage_locations"),
        ("📦  Items in Storage", "show_inventory_by_location"),
        ("📥  Purchase Orders", "show_purchase_orders"),
        ("📦  Purchased Items", "show_purchased_items"),
        ("📤  Sales Orders", "show_sales_orders"),
        ("🛒  Sold Items", "show_sold_items"),
        ("⚠️  Low Stock", "show_low_stock"),
        ("🚫  Out of Stock", "show_out_of_stock"),
        ("📥  Pending Purchases", "show_pending_purchases"),
        ("📤  Pending Sales", "show_pending_sales"),
        ("📜  Product History", "show_product_history"),
        ("📋  Transaction Log", "show_transaction_log"),
        ("📝  Log", "show_audit_log"),
    ]

    def __init__(self, user):
        super().__init__(themename="pulse")
        self.dark_mode = False
        self.current_page = "show_dashboard"
        self.user = user
        self.title(f"Inventory Management System - {user['username']} ({user['role']})")
        self.geometry("1250x720")
        self.minsize(1050, 600)
        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, minsize=230)
        self.grid_columnconfigure(1, weight=1)

        self.container = tb.Frame(self, padding=25)
        self.container.grid(row=0, column=1, sticky="nsew")
        self.container.grid_rowconfigure(0, weight=1)
        self.container.grid_columnconfigure(0, weight=1)

        sidebar = tb.Frame(self, bootstyle="primary", width=230)
        sidebar.grid(row=0, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(0, weight=1)
        sidebar.grid_columnconfigure(0, weight=1)

        nav_canvas = tb.Canvas(
             sidebar, highlightthickness=0, borderwidth=0,
                background=self.style.colors.primary
            )
        self.nav_canvas = nav_canvas

        nav_scrollbar = ttk.Scrollbar(
            sidebar, orient=VERTICAL, command=nav_canvas.yview
        )
        nav_canvas.configure(yscrollcommand=nav_scrollbar.set)
        nav_canvas.grid(row=0, column=0, sticky="nsew")

        nav_frame = tb.Frame(nav_canvas, bootstyle="primary", padding=(0, 0, 0, 8))
        nav_window = nav_canvas.create_window(
            (0, 0), window=nav_frame, anchor="nw"
        )
        tb.Label(
            nav_frame, text="📦 INVENTORY", font=("Segoe UI", 15, "bold"),
            bootstyle="inverse-primary"
        ).pack(fill=X, pady=(24, 25))
        def update_navigation(event=None):
            content_height = nav_frame.winfo_reqheight()
            viewport_height = nav_canvas.winfo_height()
            nav_canvas.itemconfigure(
                nav_window,
                height=max(content_height, viewport_height)
            )
            nav_canvas.configure(scrollregion=nav_canvas.bbox("all"))
            if nav_canvas.yview() == (0.0, 1.0):
                nav_canvas.yview_moveto(0)

        nav_frame.bind("<Configure>", update_navigation)
        nav_canvas.bind(
            "<Configure>",
            lambda event: (
                nav_canvas.itemconfigure(nav_window, width=event.width),
                nav_scrollbar.grid(row=0, column=1, sticky="ns"),
                update_navigation()
            )
        )
        nav_scrollbar.grid(row=0, column=1, sticky="ns")
        def scroll_navigation(event):
            widget = self.winfo_containing(event.x_root, event.y_root)
            while widget is not None:
                if widget == sidebar:
                    event_num = getattr(event, "num", None)
                    if event_num == 4:
                        amount = -1
                    elif event_num == 5:
                        amount = 1
                    else:
                        amount = int(-event.delta / 120)
                    nav_canvas.yview_scroll(amount, "units")
                    return "break"
                widget = widget.master

        self.bind_all("<MouseWheel>", scroll_navigation)
        self.bind_all("<Button-4>", scroll_navigation)
        self.bind_all("<Button-5>", scroll_navigation)
        self.nav_buttons = {}
        for label, method_name in self.NAV_ITEMS:
            if method_name in (
                "show_product_history", "show_transaction_log", "show_audit_log"
            ) and self.user["role"] != "admin":
                continue
            btn = tb.Button(
                nav_frame, text=label, bootstyle="dark", width=24,
                command=lambda m=method_name: self.navigate(m)
            )
            btn.pack(fill=X, pady=2, padx=8, ipady=5)
            self.nav_buttons[method_name] = btn

        divider = tk.Frame(nav_frame, height=2)
        divider.configure(bg="#8b6fd6")
        divider.pack(fill=X, padx=15, pady=15)

        self.theme_btn = tb.Button(
            nav_frame, text="🌙  Dark Mode", bootstyle="light-outline", width=20,
            command=self.toggle_theme
            
        )
        self.theme_btn.pack(pady=(0, 8), padx=8, ipady=6)

        tb.Button(
            nav_frame, text="⏻  Logout", bootstyle="danger", width=20,
            command=self.logout
        ).pack(pady=(0, 8), padx=8, ipady=6)

        self.navigate("show_dashboard")

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode
        self.style.theme_use("darkly" if self.dark_mode else "pulse")
        self.theme_btn.configure(
            text="☀️  Light Mode" if self.dark_mode else "🌙  Dark Mode"
        )
        self.nav_canvas.configure(background=self.style.colors.primary)
        self._restore_dashboard_panel_colors()
        self._apply_tree_style()
        for tree in self._find_trees(self.container):
            tree.tag_configure("hover", background=self.style.colors.inputbg)


    def logout(self):
        self.destroy()
        LoginWindow().mainloop()

    def navigate(self, method_name):
        self.current_page = method_name

        for name, btn in self.nav_buttons.items():
               btn.configure(bootstyle="primary" if name != method_name else "light")
        getattr(self, method_name)()
        self._add_export_button(method_name)

    def _find_trees(self, widget):
        found = []
        if isinstance(widget, ttk.Treeview):
            found.append(widget)
        for child in widget.winfo_children():
            found.extend(self._find_trees(child))
        return found

    def _add_export_button(self, method_name):
        if method_name in ("show_dashboard", "show_transaction_log"):
            return
        trees = self._find_trees(self.container)
        if not trees:
            return

        if method_name == "show_transaction_log" and len(trees) >= 2:
            export_target = trees[:2]
            title = "Transaction Log"
        else:
            export_target = trees[0]
            title = method_name.replace("show_", "").title()

        btn = tb.Button(
            self.container, text="📥 Export to Excel", bootstyle="success-outline",
            command=lambda: self.export_tree_to_excel(export_target, title)
        )
        btn.place(relx=1.0, y=0, anchor="ne")
        btn.lift()

    def clear_container(self):
        for widget in self.container.winfo_children():
            self._unbind_tree_events(widget)
            widget.destroy()
        self._dashboard_color_targets = []

    def _restore_dashboard_panel_colors(self):
        for widget, color in getattr(self, "_dashboard_color_targets", []):
            if widget.winfo_exists():
                widget.configure(bg=color)
                if isinstance(widget, tk.Label):
                    widget.configure(fg="white")

    def _unbind_tree_events(self, widget):
        if isinstance(widget, ttk.Treeview):
            for sequence in ("<<TreeviewSelect>>", "<Double-1>"):
                try:
                    widget.unbind(sequence)
                except TclError:
                    pass
        for child in widget.winfo_children():
            self._unbind_tree_events(child)

    def page_header(self, title, subtitle=""):
        tb.Label(self.container, text=title, font=FONT_TITLE).pack(anchor=W)
        tb.Label(self.container, text=subtitle, font=FONT_NORMAL,
                 bootstyle="secondary").pack(anchor=W, pady=(0, 20))


    def _apply_tree_style(self):
            colors = self.style.colors
            style = ttk.Style()
            style.configure("Treeview", rowheight=32, font=("Segoe UI", 10),
                            background=colors.bg, fieldbackground=colors.bg,
                            foreground=colors.fg)
            style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"),
                            background=colors.primary, foreground="#ffffff",
                            relief="flat", padding=(8, 8))
            style.map("Treeview.Heading", background=[("active", colors.primary)])

    def styled_tree(self, parent, cols, widths=None):
        self._apply_tree_style()
        colors = self.style.colors
        wrapper = tb.Frame(parent)
        wrapper.pack(fill=BOTH, expand=True)
        tree = ttk.Treeview(wrapper, columns=cols, show="headings", height=5)
        for i, c in enumerate(cols):
            tree.heading(c, text=c, anchor="center")
            tree.column(c, width=(widths[i] if widths else 120), anchor="center", stretch=True)

        tree.tag_configure("hover", background=colors.inputbg)

        vsb = ttk.Scrollbar(wrapper, orient=VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.pack(side=LEFT, fill=BOTH, expand=True)
        vsb.pack(side=RIGHT, fill=Y)

        self._bind_hover(tree)
        return tree

    def _bind_hover(self, tree):
        def on_select(event):
            try:
                if not tree.winfo_exists():
                    return
                selected = tree.selection()
                for item in tree.get_children():
                    tags = [t for t in tree.item(item, "tags") if t != "hover"]
                    if item in selected:
                        tags.append("hover")
                    tree.item(item, tags=tuple(tags))
            except TclError:
                return

        tree.bind("<<TreeviewSelect>>", on_select)

    def toolbar(self, on_search, on_add, add_label="+ Add"):
        top = tb.Frame(self.container)
        top.pack(fill=X, pady=(0, 8))
        search_entry = tb.Entry(top, width=35, bootstyle="primary")
        search_entry.pack(side=LEFT, ipady=4)
        search_entry.bind("<Return>", lambda e: on_search(search_entry.get()))
        tb.Button(top, text="🔍 Search", bootstyle="info-outline",
                  command=lambda: on_search(search_entry.get())).pack(side=LEFT, padx=8)
        tb.Button(top, text=add_label, bootstyle=SUCCESS,
                  command=on_add).pack(side=RIGHT, ipadx=10)

    def export_tree_to_excel(self, tree_or_trees, title):
            trees = tree_or_trees if isinstance(tree_or_trees, (list, tuple)) else [tree_or_trees]
            if not trees:
                messagebox.showinfo("Nothing to Export", "There are no rows to export.")
                return

            export_rows = []
            for tree in trees:
                rows = [tree.item(i)["values"] for i in tree.get_children()]
                export_rows.append((tree, rows))

            total_rows = sum(len(rows) for _, rows in export_rows)
            if total_rows == 0:
                messagebox.showinfo("Nothing to Export", "There are no rows to export.")
                return

            path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=f"{title}_{datetime.now():%Y%m%d_%H%M}.xlsx"
            )
            if not path:
                return

            wb = Workbook()
            for index, (tree, rows) in enumerate(export_rows):
                if index == 0:
                    ws = wb.active
                else:
                    ws = wb.create_sheet()

                sheet_title = title if index == 0 else (
                    "Transfer Log" if any("transfer" in str(column).lower() for column in tree["columns"]) else f"{title} {index + 1}"
                )
                ws.title = sheet_title[:31]

                columns = list(tree["columns"])
                ws.append(columns)
                for row in rows:
                    ws.append(row)

                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="5B21B6")
                    cell.alignment = Alignment(horizontal="center")

                for col in ws.columns:
                    longest = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                    ws.column_dimensions[col[0].column_letter].width = min(longest + 4, 50)

            try:
                wb.save(path)
            except PermissionError:
                messagebox.showerror(
                    "Cannot Save",
                    "The file is open in Excel. Close it and try again."
                )
                return

            try:
                db.record_export(self.user["user_id"], title, path, total_rows)
                messagebox.showinfo("Export Complete", f"Saved {total_rows} rows to:\n{path}")
            except Exception as exc:
                messagebox.showwarning(
                    "Exported",
                    f"Saved {total_rows} rows to:\n{path}\n\nThe file export succeeded, but the audit log could not be updated: {exc}"
                )

    def form_dialog(self, title, field_labels, on_save, prefill=None):
        form = tb.Toplevel(self)
        form.title(title)
        form.geometry("420x" + str(180 + 85 * len(field_labels)))
        form.resizable(False, False)
        pad = tb.Frame(form, padding=25)
        pad.pack(fill=BOTH, expand
                 =True)
        tb.Label(pad, text=title, font=FONT_HEADING).pack(anchor=W, pady=(0, 15))

        entries = {}
        for label in field_labels:
            tb.Label(pad, text=label, font=FONT_NORMAL).pack(anchor=W, pady=(8, 2))
            e = tb.Entry(pad, bootstyle="primary")
            e.pack(fill=X, ipady=5)
            if prefill and label in prefill:
                e.insert(0, prefill[label])
            entries[label] = e

        def save():
            values = {k: v.get().strip() for k, v in entries.items()}
            try:
                on_save(values)
                form.destroy()
            except Exception as e:
                messagebox.showerror("Error", str(e))

        tb.Button(pad, text="Save", bootstyle=SUCCESS, command=save).pack(
            fill=X, pady=(20, 0), ipady=6)

    def show_dashboard(self):
        self.clear_container()
        self.page_header("Dashboard", "Overview of your inventory")
        stats = db.dashboard_stats()
        fin = db.dashboard_extra()

        cards_frame = tb.Frame(self.container)
        cards_frame.pack(fill=X, pady=(0, 20))
        for col in range(6):
            cards_frame.grid_columnconfigure(col, weight=1, uniform="cards")

        def stat_card(col, icon, title, value, style, on_click):
            card = tb.Frame(cards_frame, bootstyle=style, padding=(12, 18))
            card.grid(row=0, column=col, sticky="nsew", padx=6)

            icon_lbl = tb.Label(card, text=icon, font=("Segoe UI", 16), bootstyle=f"inverse-{style}")
            icon_lbl.pack(anchor=W)
            value_lbl = tb.Label(card, text=str(value), font=("Segoe UI", 18, "bold"),
                                  bootstyle=f"inverse-{style}")
            value_lbl.pack(anchor=W, pady=(5, 0))
            title_lbl = tb.Label(card, text=title, font=("Segoe UI", 9), bootstyle=f"inverse-{style}",
                                  wraplength=110, justify=LEFT)
            title_lbl.pack(anchor=W)

            card.configure(cursor="hand2")
            for widget in (card, icon_lbl, value_lbl, title_lbl):
                widget.bind("<Button-1>", lambda e: self.navigate(on_click))

        stat_card(0, "📦", "Total Products", stats["total_products"], "info", "show_products")
        stat_card(1, "📍", "Storage Locations", fin["locations"], "success", "show_storage_locations")
        stat_card(2, "⚠️", "Low Stock", stats["low_stock"], "warning", "show_low_stock")
        stat_card(3, "🚫", "Out of Stock", stats["out_of_stock"], "danger", "show_out_of_stock")
        stat_card(4, "📥", "Pending Purchases", stats["pending_po"], "secondary", "show_pending_purchases")
        stat_card(5, "📤", "Pending Sales", stats["pending_so"], "primary", "show_pending_sales")

        panels = tb.Frame(self.container)
        panels.pack(fill=X, pady=(0, 15))
        panels.grid_columnconfigure(0, weight=1, uniform="panel")
        panels.grid_columnconfigure(1, weight=1, uniform="panel")
        self._dashboard_color_targets = []

        def panel(col, icon, title, rows, color, accent):
            box = tk.Frame(panels, padx=20, pady=16)
            box.configure(bg=color)                      # set AFTER creation
            self._dashboard_color_targets.append((box, color))
            box.grid(row=0, column=col, sticky="nsew", padx=6)
            box.grid_columnconfigure(0, weight=1)

            def make_label(text, font, row, column, sticky, colspan=1, pady=4):
                lbl = tk.Label(box, text=text, font=font)
                lbl.configure(bg=color, fg="white")      # set AFTER creation
                self._dashboard_color_targets.append((lbl, color))
                lbl.grid(row=row, column=column, columnspan=colspan,
                         sticky=sticky, pady=pady)

            make_label(f"{icon}  {title}", ("Segoe UI", 12, "bold"),
                       0, 0, "w", colspan=2, pady=0)

            line = tk.Frame(box, height=1)
            line.configure(bg=accent)
            self._dashboard_color_targets.append((line, accent))
            line.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(10, 8))

            for i, (label, value) in enumerate(rows, start=2):
                is_total = (i == len(rows) + 1)
                make_label(label,
                           ("Segoe UI", 11, "bold" if is_total else "normal"),
                           i, 0, "w")
                make_label(value,
                           ("Segoe UI", 15 if is_total else 12, "bold"),
                           i, 1, "e")

        money = lambda v: f"₹{v:,.0f}"
        panel(0, "💰", "Inventory Valuation", [
            ("Total Cost Basis", money(fin["cost_basis"])),
            ("Total Retail Value", money(fin["retail_value"])),
            ("Unrealized Gross Margin",
             f"{money(fin['gross_margin'])}  ({fin['margin_pct']:.1f}%)"),
        ], color="#db2777", accent="#f9a8d4")      # rose
        panel(1, "📈", "Sales Profitability (Completed Orders)", [
            ("Net Sales Revenue", money(fin["revenue"])),
            ("Cost of Goods Sold (COGS)", money(fin["cogs"])),
            ("Realized Gross Profit",
             f"{money(fin['gross_profit'])}  ({fin['profit_pct']:.1f}%)"),
        ], color="#4f46e5", accent="#a5b4fc")      # indigo

        tb.Label(self.container, text="All Products", font=FONT_HEADING).pack(anchor=W, pady=(5, 8))
        tree = self.styled_tree(
            self.container,
            ("ID", "Name", "Category", "Cost", "Price", "Qty", "Alert Level"),
            widths=[50, 200, 130, 90, 90, 70, 100]
        )
        for p in db.get_products():
            low = p["quantity_in_stock"] <= (p["alert_level"] or 0)
            tree.insert("", END, values=(
                p["product_id"], p["name"], p["category"], p["unit_cost"],
                p["unit_price"], p["quantity_in_stock"], p["alert_level"]
            ), tags=("low",) if low else ())
        tree.tag_configure("low", foreground="#d64545")

    def show_products(self):
        self.clear_container()
        self.page_header("Products", "Manage your product catalog")

        tree = self.styled_tree(
            self.container,
            ("ID", "Name", "Category", "Cost", "Price", "Qty", "Alert Level"),
            widths=[50, 200, 130, 90, 90, 70, 100]
        )

        def refresh(search=""):
            tree.delete(*tree.get_children())
            for p in db.get_products(search):
                low = p["quantity_in_stock"] <= (p["alert_level"] or 0)
                tree.insert("", END, values=(
                    p["product_id"], p["name"], p["category"], p["unit_cost"],
                    p["unit_price"], p["quantity_in_stock"], p["alert_level"]
                ), tags=("low",) if low else ())
            tree.tag_configure("low", foreground="#d64545")

        def add_product():
            def save(v):
                db.add_product(v["Name"], v["Category"], float(v["Unit Cost"]),
                                float(v["Unit Price"]), int(v["Alert Level"]),
                                self.user["user_id"])
                refresh()
            self.form_dialog("Add Product",
                              ["Name", "Category", "Unit Cost", "Unit Price", "Alert Level"],
                              save)

        self.toolbar(refresh, add_product, "+ Add Product")

        def edit_selected():
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            product_id = values[0]

            def save(v):
                db.update_product(product_id, v["Name"], v["Category"],
                                   float(v["Unit Cost"]), float(v["Unit Price"]),
                                    int(v["Alert Level"]),
                                    self.user["user_id"])
                refresh()

            self.form_dialog(
                "Edit Product",
                ["Name", "Category", "Unit Cost", "Unit Price", "Alert Level"],
                save,
                prefill={"Name": values[1], "Category": values[2], "Unit Cost": values[3],
                         "Unit Price": values[4], "Alert Level": values[6]}
            )

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            product_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this product?"):
                try:
                    db.delete_product(product_id, self.user["user_id"])
                    refresh()
                except Exception as e:
                    if "foreign key constraint" in str(e).lower():
                        messagebox.showerror(
                            "Cannot Delete",
                            "This product can't be deleted because it has existing "
                            "purchase or sales orders linked to it.\n\n"
                            "Delete those orders first, or keep this product for record-keeping."
                        )
                    else:
                        messagebox.showerror("Error", str(e))

        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)
        tb.Button(btns, text="✎ Edit Selected", bootstyle="info-outline",
              command=edit_selected).pack(side=LEFT, padx=8)
        

        refresh()

    def show_suppliers(self):
        self.clear_container()
        self.page_header("Suppliers", "Manage supplier records")

        tree = self.styled_tree(
            self.container, ("ID", "Name", "Category", "Contact Person", "Email", "Phone"),
            widths=[50, 180, 120, 160, 180, 120]
        )

        def refresh(search=""):
            tree.delete(*tree.get_children())
            for s in db.get_suppliers(search):
                tree.insert("", END, values=(
                    s["supplier_id"], s["name"], s["category"],
                    s["contact_person"], s["email"], s["phone"]
                ))

        def add_supplier():
            def save(v):
                db.add_supplier(v["Name"], v["Category"], v["Contact Person"], v["Email"],
                                v["Phone"], self.user["user_id"])
                refresh()
            self.form_dialog("Add Supplier", ["Name", "Category", "Contact Person", "Email", "Phone"], save)

        self.toolbar(refresh, add_supplier, "+ Add Supplier")

        def edit_selected():
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            supplier_id = values[0]

            def save(v):
                db.update_supplier(supplier_id, v["Name"], v["Category"],
                                    v["Contact Person"], v["Email"], v["Phone"],
                                    self.user["user_id"])
                refresh()

            self.form_dialog(
                "Edit Supplier",
                ["Name", "Category", "Contact Person", "Email", "Phone"],
                save,
                prefill={"Name": values[1], "Category": values[2],
                        "Contact Person": values[3], "Email": values[4], "Phone": values[5]}
            )

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            supplier_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this supplier?"):
                db.delete_supplier(supplier_id, self.user["user_id"])
                refresh()

        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)
        tb.Button(btns, text="✎ Edit Selected", bootstyle="info-outline",
              command=edit_selected).pack(side=LEFT, padx=8)
        refresh()

    def show_customers(self):
        self.clear_container()
        self.page_header("Customers", "Manage customer records")

        tree = self.styled_tree(
            self.container, ("ID", "Name", "Email", "Phone", "Address"),
            widths=[50, 180, 180, 120, 220]
        )

        def refresh(search=""):
            tree.delete(*tree.get_children())
            for c in db.get_customers(search):
                tree.insert("", END, values=(
                    c["customer_id"], c["name"], c["email"], c["phone"], c["address"]
                ))

        def add_customer():
            def save(v):
                db.add_customer(v["Name"], v["Email"], v["Phone"], v["Address"],
                                self.user["user_id"])
                refresh()
            self.form_dialog("Add Customer", ["Name", "Email", "Phone", "Address"], save)

        self.toolbar(refresh, add_customer, "+ Add Customer")

        def show_purchase_history(customer_id, customer_name):
            popup = tb.Toplevel(self)
            popup.title(f"Purchases - {customer_name}")
            popup.geometry("520x420")
            pad = tb.Frame(popup, padding=20)
            pad.pack(fill=BOTH, expand=True)
            tb.Label(pad, text=f"Products bought by {customer_name}",
                    font=FONT_HEADING).pack(anchor=W, pady=(0, 15))

            purchases = db.get_customer_purchases(customer_id)
            if not purchases:
                tb.Label(pad, text="No purchases yet.", font=FONT_NORMAL,
                        bootstyle="secondary").pack(pady=20)
                return

            ptree = self.styled_tree(pad, ("Product", "Qty", "Date", "Status"),
                                    widths=[180, 70, 110, 100])
            for row in purchases:
                ptree.insert("", END, values=(
                    row["product_name"], row["quantity"], row["order_date"], row["status"]
                ))

        def on_double_click(event):
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            show_purchase_history(values[0], values[1])

        tree.bind("<Double-1>", on_double_click)

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            customer_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this customer?"):
                db.delete_customer(customer_id, self.user["user_id"])
                refresh()

        def edit_selected():
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            customer_id = values[0]

            def save(v):
                db.update_customer(
                    customer_id, v["Name"], v["Email"], v["Phone"], v["Address"],
                    self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Edit Customer", ["Name", "Email", "Phone", "Address"], save,
                prefill={"Name": values[1], "Email": values[2],
                         "Phone": values[3], "Address": values[4]}
            )

        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)
        tb.Button(btns, text="✎ Edit Selected", bootstyle="info-outline",
                  command=edit_selected).pack(side=LEFT, padx=8)
        refresh()

    def show_storage_locations(self):
        self.clear_container()
        self.page_header("Storage Locations", "Manage stock storage locations")

        tree = self.styled_tree(
            self.container, ("ID", "Address", "Total Stock"),
            widths=[70, 480, 140]
        )

        def refresh(search=""):
            tree.delete(*tree.get_children())
            for location in db.get_storage_locations(search):
                tree.insert("", END, values=(
                    location["location_id"], location["full_address"],
                    location["total_stock"]
                ))

        def add_location():
            def save(values):
                db.add_storage_location(values["Full Address"], self.user["user_id"])
                refresh()
            self.form_dialog("Add Storage Location", ["Full Address"], save)

        self.toolbar(refresh, add_location, "+ Add Location")

        def show_location_inventory(location_id, address):
            popup = tb.Toplevel(self)
            popup.title(f"Inventory - {address}")
            popup.geometry("520x420")

            pad = tb.Frame(popup, padding=20)
            pad.pack(fill=BOTH, expand=True)
            tb.Label(
                pad,
                text=f"Products stored at {address}",
                font=FONT_HEADING
            ).pack(anchor=W, pady=(0, 15))

            inventory = db.get_location_inventory(location_id)
            if not inventory:
                tb.Label(
                    pad,
                    text="No products stored here.",
                    font=FONT_NORMAL,
                    bootstyle="secondary"
                ).pack(pady=20)
                return

            inventory_tree = self.styled_tree(
                pad,
                ("Product", "Quantity"),
                widths=[300, 120]
            )
            for item in inventory:
                inventory_tree.insert(
                    "", END,
                    values=(item["product_name"], item["stock"])
                )

        def on_double_click(event):
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            show_location_inventory(values[0], values[1])

        tree.bind("<Double-1>", on_double_click)

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            location_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this storage location?"):
                try:
                    db.delete_storage_location(location_id, self.user["user_id"])
                    refresh()
                except Exception as e:
                    if "foreign key constraint" in str(e).lower():
                        messagebox.showerror(
                            "Cannot Delete",
                            "This location can't be deleted because it has existing orders linked to it."
                        )
                    else:
                        messagebox.showerror("Error", str(e))

        def edit_selected():
            selected = tree.focus()
            if not selected:
                return
            values = tree.item(selected)["values"]
            location_id = values[0]

            def save(form_values):
                db.update_storage_location(
                    location_id, form_values["Full Address"], self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Edit Storage Location", ["Full Address"], save,
                prefill={"Full Address": values[1]}
            )

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))
        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)
        tb.Button(btns, text="✎ Edit Selected", bootstyle="info-outline",
                  command=edit_selected).pack(side=LEFT, padx=8)
        refresh()

    def show_inventory_by_location(self):
        self.clear_container()
        self.page_header("Items in Storage Locations", "Manage products stored at each location")

        tree = self.styled_tree(
            self.container,
            ("Inventory ID", "Product", "Storage Location", "Stock"),
            widths=[100, 220, 360, 100]
        )

        def refresh(search=""):
            tree.delete(*tree.get_children())
            for item in db.get_inventory_by_location(search):
                tree.insert("", END, values=(
                    item["inventory_id"], item["product_name"],
                    item["location_address"], item["stock"]
                ))

        def add_item():
            def save(values):
                db.add_inventory_by_location(
                    int(values["Product ID"]),
                    int(values["Location ID"]),
                    self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Add Item to Storage Location",
                ["Product ID", "Location ID"],
                save
            )

        self.toolbar(refresh, add_item, "+ Add Item")

        def transfer_stock():
            def save(values):
                db.transfer_stock(
                    int(values["Inventory ID"]),
                    int(values["Quantity"]),
                    int(values["Current Location"]),
                    int(values["New Location"]),
                    self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Transfer Stock",
                ["Inventory ID", "Quantity", "Current Location", "New Location"],
                save
            )

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return

            inventory_id, product_name, location_address, stock = tree.item(selected)["values"]
            if stock != 0:
                messagebox.showerror(
                    "Cannot Delete",
                    f"'{product_name}' at '{location_address}' still has stock ({stock}). "
                    "Only entries with stock 0 can be deleted."
                )
                return

            if messagebox.askyesno(
                "Confirm",
                f"Delete '{product_name}' from '{location_address}'?"
            ):
                try:
                    db.delete_inventory_by_location(inventory_id, self.user["user_id"])
                    refresh()
                except Exception as error:
                    if "foreign key constraint" in str(error).lower():
                        messagebox.showerror(
                            "Cannot Delete",
                            "This inventory entry cannot be deleted because it is referenced "
                            "by an existing purchase or sales order."
                        )
                    else:
                        messagebox.showerror("Unable to delete inventory entry", str(error))

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))
        tb.Button(
            btns,
            text="🗑 Delete Selected",
            bootstyle="danger-outline",
            command=delete_selected
        ).pack(side=LEFT)
        tb.Button(
            btns,
            text="⇄ Transfer Stock",
            bootstyle="primary-outline",
            command=transfer_stock
        ).pack(side=LEFT, padx=8)
        refresh()

    def show_purchase_orders(self):
        self.clear_container()
        self.page_header("Purchase Orders", "Stock coming IN from suppliers")

        tree = self.styled_tree(
            self.container,
            ("PO ID", "Supplier", "Date", "Status", "Note"),
            widths=[60, 180, 100, 100, 300]
        )

        def refresh():
            tree.delete(*tree.get_children())
            seen_po_ids = set()
            for po in db.get_purchase_orders():
                if po["po_id"] in seen_po_ids:
                    continue
                seen_po_ids.add(po["po_id"])
                tree.insert("", END, values=(
                    po["po_id"], po["supplier_name"], po["order_date"],
                    po["status"], po["notes"] or ""
                ))

        def add_po():
            def save(v):
                db.add_purchase_order(
                    int(v["Supplier ID"]), self.user["user_id"], v["Note"]
                )
                refresh()
            self.form_dialog("New Purchase Order",
                              ["Supplier ID", "Note"], save)

        top = tb.Frame(self.container)
        top.pack(fill=X, pady=(0, 15))
        tb.Button(top, text="+ New Purchase Order", bootstyle=SUCCESS,
                  command=add_po).pack(side=RIGHT, ipadx=10)
        tb.Label(top, text="Tip: check Products, Suppliers and Storage Locations pages for IDs",
                  font=("Segoe UI", 9), bootstyle="secondary").pack(side=LEFT)

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))

        def edit_note():
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            values = tree.item(selected)["values"]

            def save(note_values):
                db.update_purchase_order_note(
                    values[0], note_values["Note"], self.user["user_id"]
                )
                refresh()

            self.form_dialog("Update Purchase Order Note", ["Note"], save,
                             prefill={"Note": values[4]})

        def mark_status(status):
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            if not selected:
                return
            po_id = tree.item(selected)["values"][0]
            try:
                db.update_po_status(po_id, status, self.user["user_id"])
                refresh()
            except Exception as error:
                messagebox.showerror("Unable to update purchase order", str(error))

        def delete_selected():
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            if not selected:
                return
            po_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this purchase order?"):
                try:
                    db.delete_purchase_order(po_id, self.user["user_id"])
                    refresh()
                except Exception as error:
                    messagebox.showerror("Unable to delete purchase order", str(error))

        tb.Button(btns, text="✔ Mark Completed", bootstyle="success-outline",
                  command=lambda: mark_status("Completed")).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="✖ Mark Cancelled", bootstyle="warning-outline",
                  command=lambda: mark_status("Cancelled")).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="+ Add/Update Note", bootstyle="info-outline",
              command=edit_note).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)

        refresh()

    def show_sales_orders(self):
        self.clear_container()
        self.page_header("Sales Orders", "Stock going OUT to customers")

        tree = self.styled_tree(
            self.container,
            ("Order ID", "Customer", "Delivery Address", "Date", "Status", "Note"),
            widths=[70, 160, 220, 100, 100, 260]
        )

        def refresh():
            tree.delete(*tree.get_children())
            seen_order_ids = set()
            for so in db.get_sales_orders():
                if so["order_id"] in seen_order_ids:
                    continue
                seen_order_ids.add(so["order_id"])
                tree.insert("", END, values=(
                    so["order_id"], so["customer_name"], so["delivery_address"],
                    so["order_date"], so["status"], so["notes"] or ""
                ))

        def add_so():
            def save(v):
                db.add_sales_order(
                    int(v["Customer ID"]),
                    v["Delivery Address"],
                    self.user["user_id"],
                    v["Note"]
                )
                refresh()
            self.form_dialog(
                "New Sales Order",
                ["Customer ID", "Delivery Address", "Note"],
                save
            )

        top = tb.Frame(self.container)
        top.pack(fill=X, pady=(0, 15))
        tb.Button(top, text="+ New Sales Order", bootstyle=SUCCESS,
                  command=add_so).pack(side=RIGHT, ipadx=10)
        tb.Label(top, text="Tip: check Products, Customers and Storage Locations pages for IDs",
                  font=("Segoe UI", 9), bootstyle="secondary").pack(side=LEFT)

        btns = tb.Frame(self.container)
        btns.pack(fill=X, pady=(15, 0))

        def edit_note():
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            values = tree.item(selected)["values"]

            def save(note_values):
                db.update_sales_order_note(
                    values[0], note_values["Note"], self.user["user_id"]
                )
                refresh()

            self.form_dialog("Update Sales Order Note", ["Note"], save,
                             prefill={"Note": values[5]})

        def mark_status(status):
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            if not selected:
                return
            order_id = tree.item(selected)["values"][0]
            try:
                db.update_so_status(order_id, status, self.user["user_id"])
                refresh()
            except Exception as error:
                messagebox.showerror("Unable to update sales order", str(error))

        def delete_selected():
            selected_items = tree.selection()
            if not selected_items:
                return
            selected = selected_items[0]
            if not selected:
                return
            order_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this sales order?"):
                try:
                    db.delete_sales_order(order_id, self.user["user_id"])
                    refresh()
                except Exception as error:
                    messagebox.showerror("Unable to delete sales order", str(error))

        tb.Button(btns, text="✔ Mark Completed", bootstyle="success-outline",
                  command=lambda: mark_status("Completed")).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="✖ Mark Cancelled", bootstyle="warning-outline",
                  command=lambda: mark_status("Cancelled")).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="+ Add/Update Note", bootstyle="info-outline",
              command=edit_note).pack(side=LEFT, padx=(0, 8))
        tb.Button(btns, text="🗑 Delete Selected", bootstyle="danger-outline",
                  command=delete_selected).pack(side=LEFT)

        refresh()

    def show_purchased_items(self):
        self.clear_container()
        self.page_header("Purchased Items", "Items included in purchase orders")

        tree = self.styled_tree(
            self.container,
            ("Item ID", "PO ID", "Product", "Supply Location", "Storage Location", "Quantity", "Unit Cost"),
            widths=[80, 80, 180, 200, 220, 100, 120]
        )

        def refresh(order_id=""):
            tree.delete(*tree.get_children())
            order_id = order_id.strip()
            if order_id and not order_id.isdigit():
                messagebox.showwarning("Invalid Order ID", "Enter a numeric Order ID.")
                return
            for item in db.get_purchase_order_items(int(order_id) if order_id else None):
                tree.insert("", END, values=(
                    item["item_id"], item["po_id"], item["product_name"],
                    item["supplier_storage_location"], item["location_address"],
                    item["quantity"], item["unit_cost"]
                ))

        def add_item():
            def save(values):
                db.add_purchase_order_item(
                    int(values["PO ID"]),
                    int(values["Quantity"]),
                    values["Unit Cost"],
                    int(values["Inventory ID"]),
                    values["Supply Location"],
                    self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Add Purchased Item",
                ["PO ID", "Quantity", "Unit Cost", "Inventory ID", "Supply Location"],
                save
            )

        top = tb.Frame(self.container)
        top.pack(fill=X, pady=(0, 15))
        search_entry = tb.Entry(top, width=25, bootstyle="primary")
        search_entry.pack(side=LEFT, ipady=4)
        search_entry.insert(0, "")
        search_entry.bind("<Return>", lambda event: refresh(search_entry.get()))
        tb.Button(
            top, text="🔍 Search Order ID", bootstyle="info-outline",
            command=lambda: refresh(search_entry.get())
        ).pack(side=LEFT, padx=8)
        tb.Button(top, text="+ Add Purchased Item", bootstyle=SUCCESS,
                  command=add_item).pack(side=RIGHT, ipadx=10)
        tb.Label(
            top,
            text="Tip: check Purchase Orders and Products pages for IDs",
            font=("Segoe UI", 9), bootstyle="secondary"
        ).pack(side=LEFT)

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            item_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this purchased item?"):
                try:
                    db.delete_purchase_order_item(item_id, self.user["user_id"])
                    refresh(search_entry.get())
                except Exception as error:
                    messagebox.showerror("Unable to delete purchased item", str(error))

        tb.Button(
            self.container, text="🗑 Delete Selected Item",
            bootstyle="danger-outline", command=delete_selected
        ).pack(anchor=W, pady=(0, 10))
        refresh()

    def show_sold_items(self):
        self.clear_container()
        self.page_header("Sold Items", "Items included in sales orders")

        tree = self.styled_tree(
            self.container,
            ("Item ID", "Order ID", "Product", "Location", "Quantity", "Unit Price"),
            widths=[80, 90, 220, 220, 100, 120]
        )

        def refresh(order_id=""):
            tree.delete(*tree.get_children())
            order_id = order_id.strip()
            if order_id and not order_id.isdigit():
                messagebox.showwarning("Invalid Order ID", "Enter a numeric Order ID.")
                return
            for item in db.get_sales_order_items(int(order_id) if order_id else None):
                tree.insert("", END, values=(
                    item["item_id"], item["order_id"], item["product_name"],
                    item["location_address"], item["quantity"], item["unit_price"]
                ))

        def add_item():
            def save(values):
                db.add_sales_order_item(
                    int(values["Order ID"]),
                    int(values["Quantity"]),
                    values["Unit Price"],
                    int(values["Inventory ID"]),
                    self.user["user_id"]
                )
                refresh()

            self.form_dialog(
                "Add Sold Item",
                ["Order ID", "Quantity", "Unit Price", "Inventory ID"],
                save
            )

        top = tb.Frame(self.container)
        top.pack(fill=X, pady=(0, 15))
        search_entry = tb.Entry(top, width=25, bootstyle="primary")
        search_entry.pack(side=LEFT, ipady=4)
        search_entry.bind("<Return>", lambda event: refresh(search_entry.get()))
        tb.Button(
            top, text="🔍 Search Order ID", bootstyle="info-outline",
            command=lambda: refresh(search_entry.get())
        ).pack(side=LEFT, padx=8)
        tb.Button(top, text="+ Add Sold Item", bootstyle=SUCCESS,
                  command=add_item).pack(side=RIGHT, ipadx=10)
        tb.Label(
            top,
            text="Tip: check Sales Orders and Products pages for IDs",
            font=("Segoe UI", 9), bootstyle="secondary"
        ).pack(side=LEFT)

        def delete_selected():
            selected = tree.focus()
            if not selected:
                return
            item_id = tree.item(selected)["values"][0]
            if messagebox.askyesno("Confirm", "Delete this sold item?"):
                try:
                    db.delete_sales_order_item(item_id, self.user["user_id"])
                    refresh(search_entry.get())
                except Exception as error:
                    messagebox.showerror("Unable to delete sold item", str(error))

        tb.Button(
            self.container, text="🗑 Delete Selected Item",
            bootstyle="danger-outline", command=delete_selected
        ).pack(anchor=W, pady=(0, 10))
        refresh()

    def show_product_history(self):
        if self.user["role"] != "admin":
            messagebox.showerror(
                "Access Denied",
                "Product history is available only to administrators.",
                parent=self
            )
            self.navigate("show_dashboard")
            return

        self.clear_container()
        self.page_header(
            "Product History",
            "Product movement from supplier to storage and customer"
        )

        tree = self.styled_tree(
            self.container,
            ("ID", "Product", "Qty", "Supplier", "Source Location",
             "Storage Location", "Customer", "Delivery Address", "Order", "Status"),
            widths=[55, 150, 55, 150, 180, 180, 150, 220, 70, 95]
        )

        def refresh():
            tree.delete(*tree.get_children())
            for record in db.get_product_history():
                tree.insert("", END, values=(
                    record["transport_id"], record["product_name"], record["quantity"],
                    record["supplier_name"], record["source_supply_location"] or "",
                    record["storage_address"], record["customer_name"],
                    record["delivery_address"], record["order_id"],
                    record["delivery_status"]
                ))

        refresh()

    def show_transaction_log(self):
        if self.user["role"] != "admin":
            messagebox.showerror(
                "Access Denied",
                "The transaction log is available only to administrators.",
                parent=self
            )
            self.navigate("show_dashboard")
            return

        self.clear_container()
        self.page_header("Transaction Log", "Recorded stock movements and reversals")

        tree = self.styled_tree(
            self.container,
            ("Transaction ID", "Type", "Reference", "Reference ID", "Product",
             "Quantity", "User", "Date", "Notes"),
            widths=[100, 100, 130, 100, 200, 90, 120, 120, 260]
        )

        def refresh():
            tree.delete(*tree.get_children())
            for row in db.get_transaction_log():
                tree.insert("", END, values=(
                    row["transaction_id"], row["transaction_type"],
                    row["reference_type"], row["reference_id"],
                    row["product_name"], row["item_quantity"],
                    row["user_name"], row["date"], row["notes"] or ""
                ))

        refresh()

        export_buttons = tb.Frame(self.container)
        export_buttons.pack(fill=X, pady=(10, 0))
        tb.Button(
            export_buttons,
            text="📥 Export Transaction Log",
            bootstyle="success-outline",
            command=lambda: self.export_tree_to_excel(tree, "Transaction Log")
        ).pack(side=LEFT)

        tb.Label(
            self.container, text="Transfer Log", font=FONT_HEADING
        ).pack(anchor=W, pady=(20, 8))
        transfer_tree = self.styled_tree(
            self.container,
            ("Transfer ID", "Product", "Quantity", "From", "To", "User", "When"),
            widths=[100, 180, 90, 260, 260, 120, 150]
        )

        for row in db.get_transfer_log():
            transfer_tree.insert("", END, values=(
                row["transfer_id"], row["product_name"], row["quantity"],
                row["source_location"], row["destination_location"],
                row["user_name"], row["transferred_at"]
            ))

        tb.Button(
            self.container,
            text="📥 Export Transfer Log",
            bootstyle="success-outline",
            command=lambda: self.export_tree_to_excel(transfer_tree, "Transfer Log")
        ).pack(anchor=W, pady=(10, 0))

    def show_audit_log(self):
        if self.user["role"] != "admin":
            messagebox.showerror(
                "Access Denied",
                "The audit log is available only to administrators.",
                parent=self
            )
            self.navigate("show_dashboard")
            return

        self.clear_container()
        self.page_header(
            "Log",
            "Audit trail of changes to products, suppliers, customers, storage and inventory"
        )

        tree = self.styled_tree(
            self.container,
            ("Audit ID", "User", "Action", "Table", "Record ID", "Details", "When"),
            widths=[80, 120, 90, 180, 90, 480, 150]
        )

        def refresh():
            tree.delete(*tree.get_children())
            for row in db.get_audit_log():
                tree.insert("", END, values=(
                    row["audit_id"], row["user_name"], row["action"],
                    row["table_name"], row["record_id"], row["details"] or "",
                    row["created_at"]
                ))

        refresh()

    def show_low_stock(self):
        self.clear_container()
        self.page_header("Low Stock Alerts", "Products at or below alert level")

        tree = self.styled_tree(
            self.container, ("ID", "Name", "Category", "Qty", "Alert Level"),
            widths=[50, 220, 150, 90, 100]
        )
        rows = db.low_stock_alert()
        for p in rows:
            tree.insert("", END, values=(
                p["product_id"], p["name"], p["category"],
                p["quantity_in_stock"], p["alert_level"]
            ), tags=("low",))
        tree.tag_configure("low", foreground="#d64545")

        if not rows:
            tb.Label(self.container, text="✅ No low-stock products", font=FONT_NORMAL,
                      bootstyle="success").pack(pady=20)

    def show_pending_purchases(self):
        self.clear_container()
        self.page_header("Pending Purchase Orders", "Purchase orders awaiting completion")

        tree = self.styled_tree(
            self.container,
            ("PO ID", "Supplier", "Supplier Location", "Date", "Status"),
            widths=[70, 200, 250, 100, 100]
        )
        rows = db.get_pending_purchase_orders()
        for po in rows:
            tree.insert("", END, values=(
                po["po_id"], po["supplier_name"], po["supplier_storage_location"],
                po["order_date"], po["status"]
            ))

        if not rows:
            tb.Label(self.container, text="✅ No pending purchase orders", font=FONT_NORMAL,
                      bootstyle="success").pack(pady=20)

    def show_pending_sales(self):
        self.clear_container()
        self.page_header("Pending Sales Orders", "Sales orders awaiting completion")

        tree = self.styled_tree(
            self.container,
            ("Order ID", "Customer", "Delivery Address", "Date", "Status"),
            widths=[70, 200, 250, 100, 100]
        )
        rows = db.get_pending_sales_orders()
        for so in rows:
            tree.insert("", END, values=(
                so["order_id"], so["customer_name"], so["delivery_address"],
                so["order_date"], so["status"]
            ))

        if not rows:
            tb.Label(self.container, text="✅ No pending sales orders", font=FONT_NORMAL,
                      bootstyle="success").pack(pady=20)

    def show_out_of_stock(self):
        self.clear_container()
        self.page_header("Out of Stock", "Products with zero quantity")

        tree = self.styled_tree(
            self.container, ("ID", "Name", "Category"),
            widths=[60, 260, 180]
        )
        rows = db.out_of_stock_alert()
        for p in rows:
            tree.insert("", END, values=(p["product_id"], p["name"], p["category"]), tags=("out",))
        tree.tag_configure("out", foreground="#d64545")

        if not rows:
            tb.Label(self.container, text="✅ Nothing is out of stock", font=FONT_NORMAL,
                      bootstyle="success").pack(pady=20)


if __name__ == "__main__":
    LoginWindow().mainloop()
